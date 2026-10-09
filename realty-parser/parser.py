"""Парсер объявлений недвижимости (ЦИАН, Авито, Домклик) -> XLSX/CSV.

Примеры:
  python parser.py "https://www.cian.ru/cat.php?deal_type=sale&offer_type=flat&region=1&room1=1" --pages 3
  python parser.py URL_CIAN URL_AVITO URL_DOMCLICK --pages 2 -o flats.xlsx

Каждый URL - обычная страница поиска с уже выставленными фильтрами в браузере.
"""
import argparse
import csv
import os
import random
import re
import sys
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlencode, urlparse, urlunparse

from playwright.sync_api import sync_playwright

UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/124.0 Safari/537.36")
FIELDS = ["source", "url", "price", "title", "area_m2", "rooms", "floor", "address"]

# JS, выполняемый на странице: возвращает список сырых карточек
JS = {
    "cian": """() => [...document.querySelectorAll("article[data-name='CardComponent']")].map(c => {
        const q = s => (c.querySelector(s) || {}).innerText || '';
        const a = c.querySelector("a[href*='/sale/'], a[href*='/rent/']");
        return {url: a ? a.href : '', title: q("[data-mark='OfferTitle']"),
                price: q("[data-mark='MainPrice']"),
                address: [...c.querySelectorAll("[data-name='GeoLabel']")].map(e => e.innerText).join(', '),
                text: c.innerText};
    })""",
    "avito": """() => [...document.querySelectorAll("[data-marker='item']")].map(c => {
        const q = s => (c.querySelector(s) || {}).innerText || '';
        const a = c.querySelector("a[data-marker='item-title']");
        return {url: a ? a.href : '', title: q("[data-marker='item-title']"),
                price: q("[data-marker='item-price']"),
                address: q("[data-marker='item-address']"), text: c.innerText};
    })""",
    "domclick": """() => [...document.querySelectorAll("a[href*='/card/']")].map(a => {
        const c = a.closest('article, [data-e2e-id], li, div') || a;
        return {url: a.href, title: a.innerText, price: '', address: '', text: c.innerText};
    })""",
}


def detect_site(url):
    host = urlparse(url).netloc
    for name, key in (("cian", "cian"), ("avito", "avito"), ("domclick", "domclick")):
        if key in host:
            return name
    raise SystemExit(f"Неизвестный сайт: {host}")


def page_url(site, url, n):
    """Подставляет номер страницы в URL поиска."""
    u = urlparse(url)
    q = dict(parse_qsl(u.query))
    if n > 1:
        if site == "domclick":
            q["offset"] = str((n - 1) * 20)
        else:
            q["p"] = str(n)
    return urlunparse(u._replace(query=urlencode(q)))


def clean_url(site, url):
    """Убирает трекинговые параметры, оставляя постоянную ссылку."""
    u = urlparse(url)
    return urlunparse(u._replace(query="", fragment=""))


def pick_title(raw):
    """Заголовок вида «1-комн. квартира, 39 м², 5/6 этаж»; у ЦИАН в OfferTitle бывает реклама."""
    m = re.search(r"(?:\d+-комн\.|\d+-к\.|студия|апартаменты)[^\n]*?м²[^\n]*", raw["text"], re.I)
    return re.sub(r"\s+", " ", (m.group(0) if m else raw["title"])).strip()


def normalize(site, raw):
    text = " ".join([raw["title"], raw["text"]])
    area = re.search(r"(\d+[.,]?\d*)\s*м²", text) or re.search(r"(\d+[.,]?\d*)\s*сот", text)
    rooms = re.search(r"(\d+)-?\s*(?:к\.|комн)", text) or (
        re.search(r"(студи[яи])", text, re.I))
    floor = re.search(r"(\d+/\d+)\s*(?:эт|этаж)", text)
    price = re.sub(r"\s+", " ", raw["price"]).strip()
    if not price:
        m = re.search(r"\d[\d\s ]{4,}\s*₽", raw["text"])
        price = re.sub(r"\s+", " ", m.group(0)).strip() if m else ""
    return {
        "source": site,
        "url": clean_url(site, raw["url"]),
        "price": price,
        "title": pick_title(raw),
        "area_m2": area.group(1).replace(",", ".") if area else "",
        "rooms": rooms.group(1) if rooms else "",
        "floor": floor.group(1) if floor else "",
        "address": re.sub(r"\s+", " ", raw["address"]).strip(),
    }


def parse_proxy(raw):
    """http://логин:пароль@хост:порт -> настройки Playwright. ValueError, если строка не такая."""
    u = urlparse(raw.strip())
    try:
        port = u.port
    except ValueError:
        port = None
    if u.scheme not in ("http", "https", "socks5") or not u.hostname or not port:
        raise ValueError("AVITO_PROXY_URL должен быть вида http://логин:пароль@хост:порт")
    cfg = {"server": f"{u.scheme}://{u.hostname}:{port}"}
    if u.username:
        cfg["username"] = unquote(u.username)
        cfg["password"] = unquote(u.password or "")
    return cfg


def proxy_for(site):
    """Прокси только для Авито (AVITO_PROXY_URL); остальные сайты идут напрямую."""
    raw = os.getenv("AVITO_PROXY_URL") if site == "avito" else None
    return parse_proxy(raw) if raw else None


HEAVY = ("image", "media", "font")
# «Лёгкий» режим: карточки есть в самом HTML, скрипты и стили не нужны (в ~8 раз меньше трафика).
# Если в лёгком режиме карточек нет - страница перезагружается полностью (см. scrape).
LIGHT = HEAVY + ("script", "stylesheet", "xhr", "fetch")


# Сайты, которым нужен JS (карточки рисуются скриптами): грузим полностью, но только свои домены.
OWN_DOMAINS = {"avito": ("avito.ru", "avito.st")}
# Браузерный профиль (кэш скриптов + cookies) сохраняется между запусками: меньше трафика и блоков.
PROFILE_SITES = os.getenv("PROFILE_SITES", "avito").split(",")
PROFILE_DIR = Path(os.getenv("PROFILE_DIR", Path(__file__).parent / "state" / "browser"))
# Для сайтов с профилем блокируем через CDP, а не page.route: route отключает HTTP-кэш браузера.
BLOCK_URLS = [
    "*.jpg*", "*.jpeg*", "*.png*", "*.webp*", "*.gif*", "*.svg*", "*.ico*", "*.woff*", "*.ttf*",
    "*.mp4*", "*img.avito.st*", "*googletagmanager*", "*google-analytics*", "*doubleclick*",
    "*mc.yandex*", "*yandex.ru/ads*", "*an.yandex*", "*buzzoola*", "*mail.ru*", "*vk.com*",
    "*top-fwz1*", "*adfox*", "*criteo*", "*tiktok*", "*facebook*", "*uxfeedback*", "*hcaptcha*",
]


def rotate_proxy_ip():
    """Просит мобильный прокси сменить IP (ссылка из кабинета, секрет AVITO_PROXY_ROTATE_URL).
    Возвращает True, если запрос выполнен. В лог идёт только хост ссылки: путь содержит ключ."""
    import time
    import urllib.request
    raw = (os.getenv("AVITO_PROXY_ROTATE_URL") or "").strip()
    if not raw:
        return False
    host = urlparse(raw).hostname
    try:
        r = urllib.request.urlopen(raw, timeout=30)
        print(f"Смена IP запрошена ({host}): HTTP {r.status}, ответ: {r.read(160).decode('utf-8', 'ignore')[:100]!r}")
    except Exception as e:
        print(f"Смена IP не удалась (хост «{host}», ссылка должна начинаться с https://): {str(e)[:100]}")
        return False
    time.sleep(int(os.getenv("AVITO_ROTATE_WAIT", "20")))
    return True


def _router(mode, own=None):
    # mode["light"] переключается на лету; картинки и шрифты не грузим никогда;
    # own - если задано, запросы к чужим доменам (аналитика, реклама) режем
    def handle(route):
        blocked = LIGHT if mode["light"] else HEAVY
        host = urlparse(route.request.url).hostname or ""
        if route.request.resource_type in blocked or (own and not host.endswith(own)):
            return route.abort()
        return route.continue_()
    return handle


def _traffic(requests, breakdown=None):
    """Сколько байт реально пришло по завершённым запросам (сжатый размер, как считает прокси).
    breakdown: dict "тип домен" -> байты, чтобы видеть, что именно съедает трафик."""
    total = 0
    for r in requests:
        try:
            sz = r.sizes()
            b = sz.get("responseBodySize", 0) + sz.get("responseHeadersSize", 0)
        except Exception:
            continue
        total += b
        if breakdown is not None:
            key = f"{r.resource_type} {urlparse(r.url).hostname}"
            breakdown[key] = breakdown.get(key, 0) + b
    requests.clear()
    return total


def scrape(urls, pages, headless=True, delay=(3, 6)):
    rows, seen, failures = [], set(), []
    proxy = os.getenv("HTTPS_PROXY") or os.getenv("https_proxy")
    # не сообщать сайту, что браузером управляет программа
    launch = {"headless": headless, "args": ["--no-sandbox", "--disable-blink-features=AutomationControlled"]}
    if proxy:
        launch["proxy"] = {"server": proxy}
    if os.getenv("CHROME_CHANNEL"):  # на GitHub Actions: предустановленный google-chrome
        launch["channel"] = os.getenv("CHROME_CHANNEL")
    elif os.path.exists("/opt/pw-browsers/chromium"):
        launch["executable_path"] = "/opt/pw-browsers/chromium"
    with sync_playwright() as p:
        browser = p.chromium.launch(**launch)
        # User-Agent должен совпадать с реальной версией браузера (её же сайт видит в Client Hints),
        # иначе расхождение выдаёт автоматику
        major = browser.version.split(".")[0]
        ua = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) "
              f"Chrome/{major}.0.0.0 Safari/537.36")
        look = {"locale": "ru-RU", "user_agent": ua, "timezone_id": "Europe/Moscow",
                "viewport": {"width": 1366, "height": 768}}
        pages_by_site, modes, finished, traffic, breakdown, persistent = {}, {}, {}, {}, {}, []
        for url in urls:
            site = detect_site(url)
            if site not in pages_by_site:
                kw = {"proxy": proxy_for(site)} if proxy_for(site) else {}
                if site in PROFILE_SITES:
                    opts = {k: v for k, v in launch.items() if k != "proxy"}
                    ctx = p.chromium.launch_persistent_context(
                        str(PROFILE_DIR / site), **look, **opts, **kw)
                    persistent.append(ctx)
                else:
                    ctx = browser.new_context(**look, **kw)
                light = os.getenv("LIGHT_MODE", "1") == "1" and site not in OWN_DOMAINS
                modes[site] = {"light": light and site not in PROFILE_SITES}
                finished[site], traffic[site], breakdown[site] = [], 0, {}
                ctx.on("requestfinished", lambda r, bucket=finished[site]: bucket.append(r))
                pg = ctx.pages[0] if ctx.pages else ctx.new_page()
                if site in PROFILE_SITES:
                    cdp = ctx.new_cdp_session(pg)
                    cdp.send("Network.enable")
                    cdp.send("Network.setBlockedURLs", {"urls": BLOCK_URLS})
                else:
                    ctx.route("**/*", _router(modes[site], OWN_DOMAINS.get(site)))
                pages_by_site[site] = pg
            page = pages_by_site[site]
            for n in range(1, pages + 1):
                target = page_url(site, url, n)
                def load():
                    r = page.goto(target, timeout=45000, wait_until="domcontentloaded")
                    page.wait_for_timeout(4000)
                    page.mouse.wheel(0, 6000)  # дозагрузка ленивых карточек
                    page.wait_for_timeout(1500)
                    return r, page.evaluate(JS[site])

                try:
                    resp, items = load()
                    # 403/429 - блок IP: меняем адрес прокси (если есть ссылка) и пробуем снова
                    tries = int(os.getenv("AVITO_ROTATE_TRIES", "3")) if site == "avito" else 1
                    for attempt in range(tries):
                        code = resp.status if resp else 0
                        if items or code not in (403, 429):
                            break
                        print(f"[{site}] HTTP {code} - попытка {attempt + 1}/{tries}")
                        if site == "avito" and os.getenv("AVITO_PROXY_ROTATE_URL"):
                            if not rotate_proxy_ip():
                                break
                        else:
                            page.wait_for_timeout(int(random.uniform(8, 15) * 1000))
                        resp, items = load()
                    # полную загрузку включаем, только если страница открылась нормально, но без карточек
                    if not items and modes[site]["light"] and resp and resp.status == 200:
                        print(f"[{site}] в лёгком режиме карточек нет - перехожу на полную загрузку")
                        modes[site]["light"] = False
                        resp, items = load()
                except Exception as e:
                    print(f"[{site} p{n}] ошибка: {str(e)[:100]}", file=sys.stderr)
                    failures.append(f"{site}: {str(e)[:80]}")
                    break
                new = 0
                for raw in items:
                    if not raw["url"]:
                        continue
                    row = normalize(site, raw)
                    row["search"] = url
                    if row["url"] in seen:
                        continue
                    seen.add(row["url"])
                    rows.append(row)
                    new += 1
                status = resp.status if resp else "?"
                if not items:
                    try:
                        print(f"[{site}] страница без карточек: «{page.title()[:80]}»")
                    except Exception:
                        pass
                kb = _traffic(finished[site], breakdown[site]) / 1024
                traffic[site] += kb
                mode = " (лёгкий режим)" if modes[site]["light"] else ""
                print(f"[{site} p{n}] HTTP {status}, карточек: {len(items)}, новых: {new}, "
                      f"трафик ~{kb:.0f} КБ{mode}")
                if not items:
                    failures.append(f"{site} p{n}: HTTP {status}, карточек нет (блокировка?)")
                if new == 0:
                    if site == "domclick":
                        print("  Домклик закрыт антибот-защитой (Qrator): "
                              "попробуй --headed на своём компьютере.", file=sys.stderr)
                    break
                lo, hi = (6, 12) if site == "avito" else delay  # Авито быстро режет по частоте
                page.wait_for_timeout(int(random.uniform(lo, hi) * 1000))
        for ctx in persistent:
            ctx.close()  # сохраняет профиль на диск
        browser.close()
    for site, kb in traffic.items():
        print(f"Трафик {site} за запуск: ~{kb / 1024:.2f} МБ")
        top = sorted(breakdown[site].items(), key=lambda x: -x[1])[:8]
        for key, b in top:
            print(f"    {b / 1024:7.0f} КБ  {key}")
    return rows, failures


def save(rows, path):
    if path.endswith(".csv"):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, FIELDS, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        return
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Объявления"
    ws.append(FIELDS)
    for r in rows:
        ws.append([r[k] for k in FIELDS])
        ws.cell(ws.max_row, 2).hyperlink = r["url"]
    for col, w in zip("ABCDEFGH", (10, 60, 16, 40, 10, 8, 8, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("urls", nargs="+", help="URL страниц поиска (ЦИАН / Авито / Домклик)")
    ap.add_argument("--pages", type=int, default=1, help="сколько страниц выдачи на каждый URL")
    ap.add_argument("-o", "--output", default="listings.xlsx", help=".xlsx или .csv")
    ap.add_argument("--headed", action="store_true", help="показать окно браузера")
    a = ap.parse_args()
    rows, _ = scrape(a.urls, a.pages, headless=not a.headed)
    save(rows, a.output)
    print(f"Готово: {len(rows)} объявлений -> {a.output}")


if __name__ == "__main__":
    main()
